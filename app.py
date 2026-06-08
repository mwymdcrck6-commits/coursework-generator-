import streamlit as st
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from io import BytesIO
from datetime import datetime

st.set_page_config(page_title="Генератор курсовой", page_icon="📘", layout="centered")

st.title("📘 Генератор курсовой работы с наполнением")
st.markdown("**С примерами текста по ГОСТ**")

with st.form("course_form"):
    col1, col2 = st.columns(2)
    with col1:
        title = st.text_area("Название курсовой", "Влияние цифровых технологий на современный бизнес", height=80)
        student = st.text_input("ФИО студента", "Иванов Иван Иванович")
        group = st.text_input("Группа", "ЭК-21-1")
    with col2:
        teacher = st.text_input("ФИО руководителя", "Петров Пётр Петрович")
        university = st.text_input("Университет", "ФГБОУ ВО «Российский университет»")
        city = st.text_input("Город", "Москва")
        year = st.number_input("Год", value=datetime.now().year)

    submitted = st.form_submit_button("Создать курсовую с наполнением", type="primary")

if submitted:
    if not title or not student:
        st.error("Заполните название и ФИО")
    else:
        with st.spinner("Создаём документ с примерным наполнением..."):
            doc = Document()

            # Форматирование
            section = doc.sections[0]
            section.top_margin = Cm(2)
            section.bottom_margin = Cm(2)
            section.left_margin = Cm(3)
            section.right_margin = Cm(1.5)

            style = doc.styles['Normal']
            style.font.name = 'Times New Roman'
            style.font.size = Pt(14)
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

            def centered(text, size=14, bold=False):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(text)
                run.font.size = Pt(size)
                if bold: run.bold = True

            # Титульный лист
            centered("МИНИСТЕРСТВО НАУКИ И ВЫСШЕГО ОБРАЗОВАНИЯ РОССИЙСКОЙ ФЕДЕРАЦИИ")
            centered(university)
            doc.add_paragraph()
            centered("КАФЕДРА ИНФОРМАЦИОННЫХ ТЕХНОЛОГИЙ")
            for _ in range(6): doc.add_paragraph()
            centered(title.upper(), 16, True)
            for _ in range(8): doc.add_paragraph()
            p = doc.add_paragraph(f"Студент: {student}\nГруппа: {group}")
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p = doc.add_paragraph(f"Руководитель: {teacher}")
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            doc.add_page_break()

            # Оглавление
            centered("ОГЛАВЛЕНИЕ", 14, True)
            doc.add_paragraph("Введение ............................................... 3")
            doc.add_paragraph("1. Теоретические основы ................................ 4")
            doc.add_paragraph("2. Анализ практического применения ..................... 12")
            doc.add_paragraph("Заключение ............................................. 20")
            doc.add_paragraph("Список литературы ..................................... 22")
            doc.add_page_break()

            # Наполнение
            doc.add_heading("ВВЕДЕНИЕ", level=1)
            doc.add_paragraph("Актуальность темы исследования обусловлена быстрым развитием